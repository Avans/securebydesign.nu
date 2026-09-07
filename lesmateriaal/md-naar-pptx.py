#!/usr/bin/env python3
"""Zet de Marp-markdown van de lesdecks om in een bewerkbare PowerPoint.

Gebruikt standaard de Avans-huisstijltemplate (avans_template.pptx): de slides
komen uit de layouts van de template, dus lettertype, kleur en logo kloppen en
je kunt in PowerPoint gewoon via 'Indeling' een andere Avans-slide kiezen.

Waarom niet marp-cli: die exporteert elke slide als plaatje en laat de
docentnotities vallen. Dit script maakt echte tekstvakken, echte tabellen en
zet het HTML-commentaar onder elke slide in het notitieveld van PowerPoint.

Gebruik:
    python3 md-naar-pptx.py week1/w1b*.md -o week1/_pptx
    python3 md-naar-pptx.py week1/w1b1-*.md --template andere.pptx
    python3 md-naar-pptx.py week1/w1b1-*.md --geen-template   # kale opmaak

Vereist: python-pptx  (pip install python-pptx)
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt
except ImportError:  # pragma: no cover
    sys.exit("python-pptx ontbreekt. Installeer met: pip install python-pptx")

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE / "avans_template.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Het stramien van de Avans-layouts: titel op 0.79", tekst begint op 1.74".
MARGIN_L = Inches(0.79)
BODY_W = Inches(11.76)
BODY_TOP = Inches(1.74)
BODY_BOTTOM = Inches(6.95)

# Avans-huisstijl (accent1 uit het thema van de template).
AVANS_ROOD = RGBColor(0xC6, 0x00, 0x2A)
INK = RGBColor(0x00, 0x00, 0x00)
INK2 = RGBColor(0x4D, 0x4D, 0x4D)
MONO = "Consolas"

# Layouts uit avans_template.pptx, per rol.
LAYOUTS = {
    "cover": "Opening slide - rood - zonder foto",
    "section": "Titelslide basic rood",
    "body": "Lege dia",
    "titleonly": "Alleen titel",
}

# ---------------------------------------------------------------- parsing


def split_frontmatter(text: str) -> tuple[dict, str]:
    """Haal de Marp-frontmatter eraf. Geeft (meta, rest)."""
    if not text.startswith("---"):
        return {}, text
    end = text.find("\n---", 3)
    if end == -1:
        return {}, text
    meta = {}
    for line in text[3:end].splitlines():
        if ":" in line and not line.startswith(" "):
            k, _, v = line.partition(":")
            meta[k.strip()] = v.strip().strip("\"'")
    return meta, text[end + 4 :]


def split_slides(body: str) -> list[str]:
    """Splits op een regel die exact --- is, buiten codeblokken om."""
    slides, buf, in_code = [], [], False
    for line in body.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
        if not in_code and line.strip() == "---":
            slides.append("\n".join(buf))
            buf = []
            continue
        buf.append(line)
    slides.append("\n".join(buf))
    return [s for s in (x.strip("\n") for x in slides) if s.strip()]


COMMENT_RE = re.compile(r"<!--(.*?)-->", re.S)
DIRECTIVE_RE = re.compile(r"^\s*_?[a-zA-Z]+\s*:", re.M)


def extract_notes(slide: str) -> tuple[str, str]:
    """Haal HTML-commentaar uit de slide; dat wordt de docentnotitie.

    Marp-directives (<!-- _class: lead -->) zijn geen notitie en verdwijnen.
    """
    notes: list[str] = []

    def take(m: re.Match) -> str:
        inner = m.group(1).strip()
        if inner and not (DIRECTIVE_RE.match(inner) and len(inner.splitlines()) == 1):
            notes.append(inner)
        return ""

    return COMMENT_RE.sub(take, slide).strip("\n"), "\n\n".join(notes).strip()


def parse_blocks(content: str) -> list[tuple]:
    """Zet de slide-markdown om in blokken: (soort, data)."""
    blocks: list[tuple] = []
    lines = content.splitlines()
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(code)))
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|[\s:|-]+\|$", lines[i + 1].strip()
        ):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            head = [c.strip() for c in rows[0].strip("|").split("|")]
            body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
            blocks.append(("table", (head, body)))
            continue

        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(q for q in quote if q)))
            continue

        if re.match(r"^[-*]\s+", stripped) or re.match(r"^\d+[.)]\s+", stripped):
            items: list[tuple[str, str]] = []
            while i < len(lines):
                s = lines[i].strip()
                mb = re.match(r"^[-*]\s+(.*)$", s)
                mo = re.match(r"^(\d+)[.)]\s+(.*)$", s)
                if mb:
                    items.append(("•", mb.group(1)))
                elif mo:
                    items.append((f"{mo.group(1)}.", mo.group(2)))
                elif s and items and not re.match(r"^(#{1,4}\s|\||>|```)", s):
                    items[-1] = (items[-1][0], items[-1][1] + " " + s)
                else:
                    break
                i += 1
            blocks.append(("list", items))
            continue

        para = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,4}\s|[-*]\s|\d+[.)]\s|\||>|```)", lines[i].strip()
        ):
            para.append(lines[i].strip())
            i += 1
        if para:
            blocks.append(("p", " ".join(para)))
    return blocks


# ---------------------------------------------------------------- rendering

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_runs(paragraph, text, size=None, color=None, bold=False, italic=False):
    """Zet **vet**, *cursief* en `code` om in losse runs."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)  # links → tekst
    for part in INLINE_RE.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        f = run.font
        if size is not None:
            f.size = size
        if color is not None:
            f.color.rgb = color
        f.bold = bold
        f.italic = italic
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            f.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            f.name = MONO
            if size is not None:
                f.size = Pt(size.pt - 2)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run.text = part[1:-1]
            f.italic = True
        else:
            run.text = part


def plain(text: str) -> str:
    """Markdown-opmaak eruit — voor titels in placeholders."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return re.sub(r"[*`]", "", text)


def fit_size(text: str, steps: list[tuple[int, int]]) -> Pt:
    for max_len, pt in steps:
        if len(text) <= max_len:
            return Pt(pt)
    return Pt(steps[-1][1])


def find_layout(prs, name: str):
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name.strip() == name:
                return layout
    return None


def strip_slides(prs) -> None:
    """Gooi de voorbeeldslides uit de template; layouts blijven staan."""
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        rid = sld.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        prs.part.drop_rel(rid)
        xml_slides.remove(sld)


def tidy_placeholders(slide, used: set[int]) -> None:
    """Decoratieve placeholders krijgen een spatie, net als in de template."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx not in used and ph.has_text_frame:
            ph.text_frame.text = " "


def textbox(slide, top, height, width=BODY_W, left=MARGIN_L):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def est_height(text: str, size_pt: float, width_in: float) -> float:
    """Ruwe hoogteschatting in inches — genoeg om blokken te stapelen."""
    per_line = max(int(width_in * 96 / (size_pt * 0.52)), 10)
    lines = max(1, -(-len(text) // per_line))
    return lines * (size_pt * 1.45) / 72


def flow_blocks(slide, blocks, y, width=BODY_W, left=MARGIN_L) -> list[str]:
    """Zet blokken als losse vormen onder elkaar. Geeft de blokken die niet pasten."""
    width_in = width / Inches(1)
    rest: list[str] = []
    for kind, data in blocks:
        if y > BODY_BOTTOM - Inches(0.3):
            rest.append(kind)
            continue

        if kind in ("h2", "h3", "h4"):
            size = Pt(22 if kind == "h2" else 20 if kind == "h3" else 16)
            h = Inches(est_height(plain(data), size.pt, width_in) + 0.12)
            add_runs(textbox(slide, y, h, width, left).paragraphs[0], data, size,
                     AVANS_ROOD, bold=True)
            y += h + Inches(0.12)

        elif kind == "p":
            h = Inches(est_height(plain(data), 16, width_in) + 0.06)
            add_runs(textbox(slide, y, h, width, left).paragraphs[0], data, Pt(16), INK)
            y += h + Inches(0.14)

        elif kind == "quote":
            h = Inches(est_height(plain(data), 17, width_in - 0.35) + 0.1)
            tf = textbox(slide, y, h, width - Inches(0.35), left + Inches(0.35))
            add_runs(tf.paragraphs[0], data, Pt(17), AVANS_ROOD, italic=True)
            bar = slide.shapes.add_shape(1, left, y, Inches(0.05), h)
            bar.fill.solid()
            bar.fill.fore_color.rgb = AVANS_ROOD
            bar.line.fill.background()
            bar.shadow.inherit = False
            y += h + Inches(0.18)

        elif kind == "list":
            total = sum(est_height(plain(t), 16, width_in - 0.4) for _, t in data)
            total += 0.09 * len(data)
            tf = textbox(slide, y, Inches(total), width, left)
            for idx, (marker, item) in enumerate(data):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.space_after = Pt(6)
                run = p.add_run()
                run.text = f"{marker}  "
                run.font.size = Pt(16)
                run.font.color.rgb = AVANS_ROOD
                add_runs(p, item, Pt(16), INK)
            y += Inches(total) + Inches(0.16)

        elif kind == "code":
            lines = data.splitlines() or [""]
            h = Inches(len(lines) * 0.19 + 0.16)
            tf = textbox(slide, y, h, width, left)
            for idx, line in enumerate(lines):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                run = p.add_run()
                run.text = line
                run.font.size = Pt(12)
                run.font.name = MONO
                run.font.color.rgb = INK2
            y += h + Inches(0.16)

        elif kind == "table":
            head, rows = data
            nrows, ncols = len(rows) + 1, len(head)
            h = Inches(0.34 * nrows)
            tbl = slide.shapes.add_table(nrows, ncols, left, y, width, h).table
            for c, cell_text in enumerate(head):
                cell = tbl.cell(0, c)
                cell.text_frame.clear()
                add_runs(cell.text_frame.paragraphs[0], cell_text, Pt(13), bold=True)
            for r, row in enumerate(rows, start=1):
                for c in range(ncols):
                    cell = tbl.cell(r, c)
                    cell.text_frame.clear()
                    add_runs(
                        cell.text_frame.paragraphs[0],
                        row[c] if c < len(row) else "",
                        Pt(13),
                    )
            y += h + Inches(0.2)
    return rest


def fill_textframe(tf, blocks) -> None:
    """Vul een placeholder-tekstvak met de blokken (behoudt de themastijl)."""
    tf.clear()
    first = True

    def para():
        nonlocal first
        if first:
            first = False
            return tf.paragraphs[0]
        return tf.add_paragraph()

    for kind, data in blocks:
        if kind in ("h2", "h3", "h4"):
            p = para()
            p.space_before = Pt(6)
            p.space_after = Pt(4)
            add_runs(p, data, Pt(22 if kind == "h2" else 20), AVANS_ROOD, bold=True)
        elif kind == "p":
            p = para()
            p.space_after = Pt(8)
            add_runs(p, data, Pt(16), INK)
        elif kind == "quote":
            p = para()
            p.space_before = Pt(6)
            p.space_after = Pt(8)
            add_runs(p, data, Pt(17), AVANS_ROOD, italic=True)
        elif kind == "list":
            for marker, item in data:
                p = para()
                p.space_after = Pt(6)
                run = p.add_run()
                run.text = f"{marker}  "
                run.font.size = Pt(16)
                run.font.color.rgb = AVANS_ROOD
                add_runs(p, item, Pt(16), INK)


def render(md_path: Path, out_path: Path, template: Path | None) -> tuple:
    text = md_path.read_text(encoding="utf-8")
    meta, body = split_frontmatter(text)

    prs = Presentation(str(template)) if template else Presentation()
    if template:
        strip_slides(prs)
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    lay = {}
    if template:
        for role, name in LAYOUTS.items():
            found = find_layout(prs, name)
            if found is None:
                sys.exit(f"Layout {name!r} ontbreekt in {template.name}")
            lay[role] = found
    blank = None if template else prs.slide_layouts[6]

    n_notes, overflow = 0, []
    slides = split_slides(body)

    for n, raw in enumerate(slides, start=1):
        content, notes = extract_notes(raw)
        blocks = parse_blocks(content)
        head_i = next(
            (i for i, b in enumerate(blocks) if b[0] in ("h1", "h2")), None
        )
        heads = [blocks[head_i]] if head_i is not None else []
        # alleen de eerste kop wordt de slidetitel; latere koppen blijven in de body
        rest = [b for i, b in enumerate(blocks) if i != head_i]
        title = plain(blocks[head_i][1]) if head_i is not None else ""
        only_headings = all(k.startswith("h") for k, _ in blocks) and bool(blocks)
        has_shapes = any(k in ("table", "code") for k, _ in blocks)

        if not template:
            role = "plain"
        elif n == 1:
            role = "cover"
        elif only_headings:
            role = "section"
        elif has_shapes:
            role = "titleonly"
        else:
            role = "body"

        # -------------------------------------------------- cover
        if role == "cover":
            slide = prs.slides.add_slide(lay["cover"])
            slide.placeholders[0].text_frame.text = title
            slide.placeholders[0].text_frame.paragraphs[0].runs[0].font.size = fit_size(
                title, [(24, 44), (48, 34), (999, 28)]
            )
            tf = slide.placeholders[21].text_frame
            fill_textframe(tf, rest)
            tidy_placeholders(slide, {0, 21})

        # -------------------------------------------------- sectiedivider
        elif role == "section":
            slide = prs.slides.add_slide(lay["section"])
            tf = slide.placeholders[0].text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            add_runs(p, title, fit_size(title, [(28, 60), (48, 44), (999, 34)]))
            for kind, data in rest:
                if kind in ("h2", "h3", "h4"):
                    sub = tf.add_paragraph()
                    sub.space_before = Pt(10)
                    add_runs(sub, data, Pt(26))
            tidy_placeholders(slide, {0})

        # -------------------------------------------------- tabel/code
        elif role == "titleonly":
            slide = prs.slides.add_slide(lay["titleonly"])
            tp = slide.placeholders[0].text_frame
            tp.text = title
            tp.paragraphs[0].runs[0].font.size = fit_size(
                title, [(28, 36), (48, 30), (75, 26), (999, 22)]
            )
            left = flow_blocks(slide, rest, BODY_TOP)
            if left:
                overflow.append(n)
            tidy_placeholders(slide, {0})

        # -------------------------------------------------- gewone tekstslide
        elif role == "body":
            slide = prs.slides.add_slide(lay["body"])
            tp = slide.placeholders[0].text_frame
            tp.text = title
            tp.paragraphs[0].runs[0].font.size = fit_size(
                title, [(28, 40), (48, 32), (75, 26), (999, 22)]
            )
            fill_textframe(slide.placeholders[21].text_frame, rest)
            tidy_placeholders(slide, {0, 21})

        # -------------------------------------------------- zonder template
        else:
            slide = prs.slides.add_slide(blank)
            y = Inches(0.55)
            if heads:
                size = Pt(40) if heads[0][0] == "h1" else Pt(28)
                h = Inches(est_height(title, size.pt, BODY_W / Inches(1)) + 0.12)
                add_runs(textbox(slide, y, h).paragraphs[0], title, size, INK, bold=True)
                y += h + Inches(0.14)
            if flow_blocks(slide, rest, y):
                overflow.append(n)

        if notes:
            slide.notes_slide.notes_text_frame.text = notes
            n_notes += 1

    if meta.get("title"):
        prs.core_properties.title = meta["title"]
    if meta.get("description"):
        prs.core_properties.comments = meta["description"]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return len(slides), n_notes, sorted(set(overflow))


def main() -> None:
    ap = argparse.ArgumentParser(description="Marp-markdown → Avans-PowerPoint")
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("-o", "--outdir", type=Path, default=None)
    ap.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    ap.add_argument(
        "--geen-template", action="store_true", help="kale opmaak, zonder huisstijl"
    )
    args = ap.parse_args()

    template = None if args.geen_template else args.template
    if template and not template.exists():
        print(f"⚠ template {template} niet gevonden — kale opmaak", file=sys.stderr)
        template = None

    for md in args.files:
        out = (args.outdir or md.parent) / (md.stem + ".pptx")
        slides, notes, overflow = render(md, out, template)
        warn = f"  ⚠ vol: slide {', '.join(map(str, overflow))}" if overflow else ""
        print(f"{md.name:42s} → {out.name:42s} {slides:3d} slides, {notes:3d} notities{warn}")


if __name__ == "__main__":
    main()
